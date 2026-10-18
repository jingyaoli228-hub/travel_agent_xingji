from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"D:\travel agent2")
SOURCE = ROOT / "docs" / "行迹微信小程序第二版技术与产品说明.md"
OUTPUT = ROOT / "docs" / "行迹微信小程序第二版技术与产品说明.docx"
SCREENSHOT = ROOT / "output" / "miniprogram-final-plan.png"

NAVY = "173B66"
BLUE = "376D9C"
PALE = "EAF1F7"
LIGHT = "F6F8FA"
BORDER = "D9D9D9"
TEXT = RGBColor(35, 43, 51)
MUTED = RGBColor(91, 104, 118)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), BORDER)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Microsoft YaHei", size=10.5, bold=None, color=TEXT) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = color


def add_inline(paragraph, text: str, size=10.5, color=TEXT) -> None:
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, size=size, color=color)
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Consolas", size=max(9, size - 0.5), color=RGBColor(26, 67, 104))
        else:
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True, color=color)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=size, color=color)


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.35) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    set_run_font(run, size=9, color=MUTED)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT

    title = styles["Title"]
    title.font.name = "Microsoft YaHei"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)

    for name, size, before, after in (("Heading 1", 17, 18, 8), ("Heading 2", 13, 14, 5), ("Heading 3", 11.5, 10, 4)):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_cover(doc: Document) -> None:
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("行迹微信小程序第二版技术与产品说明")
    set_paragraph_spacing(p, after=16, line=1.0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("产品功能  技术架构  数据流程  发布边界")
    set_run_font(r, size=13, color=RGBColor(55, 85, 112))
    set_paragraph_spacing(p, after=80, line=1.0)

    meta = [
        ("版本", "2.0"),
        ("日期", "2026年10月18日"),
        ("范围", "当前代码仓库  腾讯云部署  微信开发者工具验证版本"),
        ("状态", "受控测试版"),
    ]
    table = doc.add_table(rows=len(meta), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.15)
    table.columns[1].width = Inches(4.9)
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for i, (key, value) in enumerate(meta):
        for cell in table.rows[i].cells:
            set_cell_margins(cell, top=130, bottom=130)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(table.cell(i, 0), PALE)
        p0 = table.cell(i, 0).paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(p0, key, size=10)
        p1 = table.cell(i, 1).paragraphs[0]
        add_inline(p1, value, size=10)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=85, after=0, line=1.0)
    r = p.add_run("Travel Planning Agent")
    set_run_font(r, size=10, color=MUTED)
    doc.add_page_break()


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if i == 0:
                set_cell_shading(cell, NAVY)
            elif i % 2 == 0:
                set_cell_shading(cell, LIGHT)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (j == 0 and len(row[j] if j < len(row) else "") < 24) else WD_ALIGN_PARAGRAPH.LEFT
            text = row[j].strip() if j < len(row) else ""
            add_inline(p, text, size=8.6, color=RGBColor(255, 255, 255) if i == 0 else TEXT)
            set_paragraph_spacing(p, after=0, line=1.18)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_screenshot(doc: Document) -> None:
    if not SCREENSHOT.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(SCREENSHOT), width=Inches(3.05))
    doc_pr = shape._inline.docPr
    doc_pr.set("name", "微信小程序逐日行程界面")
    doc_pr.set("descr", "微信开发者工具中显示三日旅行规划、每日活动和阶段标签的界面截图")
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = False
    r = cap.add_run("图 1  微信开发者工具中的逐日行程界面")
    set_run_font(r, size=9, color=MUTED)
    set_paragraph_spacing(cap, after=8, line=1.0)


def parse_table(lines: list[str], start: int):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        parts = [x.strip() for x in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", x or "") for x in parts):
            rows.append(parts)
        i += 1
    return rows, i


def build() -> None:
    source_lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    configure_styles(doc)
    add_cover(doc)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("行迹微信小程序第二版技术与产品说明")
    set_run_font(r, size=8.5, color=MUTED)
    footer = section.footer.paragraphs[0]
    add_page_number(footer)

    i = 1  # skip Markdown title, cover already contains it
    in_code = False
    code_lines: list[str] = []
    screenshot_added = False
    while i < len(source_lines):
        raw = source_lines[i]
        stripped = raw.strip()
        if stripped.startswith("版本 2.0") or stripped.startswith("编制日期") or stripped.startswith("适用范围"):
            i += 1
            continue
        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                in_code = False
                if code_lines:
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Inches(0.28)
                    p.paragraph_format.right_indent = Inches(0.16)
                    p.paragraph_format.keep_together = True
                    add_inline(p, "\n".join(code_lines), size=8.4, color=RGBColor(44, 61, 76))
                    for run in p.runs:
                        set_run_font(run, name="Consolas", size=8.4, color=RGBColor(44, 61, 76))
                    set_paragraph_spacing(p, before=4, after=8, line=1.1)
            i += 1
            continue
        if in_code:
            code_lines.append(raw)
            i += 1
            continue
        if not stripped:
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(source_lines) and source_lines[i + 1].strip().startswith("|"):
            rows, i = parse_table(source_lines, i)
            add_table(doc, rows)
            continue
        m = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if m:
            level = len(m.group(1)) - 1
            p = doc.add_paragraph(style=f"Heading {min(level, 3)}")
            add_inline(p, m.group(2), size={1: 17, 2: 13, 3: 11.5}.get(level, 11.5), color=RGBColor(0, 0, 0))
            if m.group(2) == "页面结构" and not screenshot_added:
                add_screenshot(doc)
                screenshot_added = True
            i += 1
            continue
        m = re.match(r"^[-*]\s+(.+)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, m.group(1), size=10.2)
            set_paragraph_spacing(p, after=3, line=1.28)
            i += 1
            continue
        m = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, m.group(2), size=10.2)
            set_paragraph_spacing(p, after=3, line=1.28)
            i += 1
            continue
        p = doc.add_paragraph()
        add_inline(p, stripped, size=10.5)
        set_paragraph_spacing(p, after=6, line=1.36)
        i += 1

    props = doc.core_properties
    props.title = "行迹微信小程序第二版技术与产品说明"
    props.subject = "产品功能 技术架构 API RAG Agent LangGraph 意图识别和记忆设计"
    props.author = "Travel Planning Agent 项目"
    props.keywords = "微信小程序, LangGraph, RAG, Agent, 旅行规划, 高德, DeepSeek"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
